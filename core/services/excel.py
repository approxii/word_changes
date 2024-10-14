from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from core.services.base import BaseDocumentService


class ExcelService(BaseDocumentService):
    def __init__(self):
        self.workbook = None

    def load(self, file) -> None:
        self.workbook = load_workbook(file)

    def update(self, params: dict) -> None:
        if not self.workbook:
            raise ValueError("Excel файл не загружен.")
        for sheet_name in self.workbook.sheetnames:
            sheet = self.workbook[sheet_name]
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.hyperlink and cell.hyperlink.display in params:
                        value = params.get(cell.hyperlink.display)

                        if cell.hyperlink.location:
                            location_sheet, location_cell = (
                                cell.hyperlink.location.split("!")
                            )
                            location_sheet = location_sheet.replace("'", "")
                            start_cell = self.workbook[location_sheet][location_cell]

                            if isinstance(
                                value, list
                            ):  # Если значение словаря - список из списков
                                for i, sublist in enumerate(
                                    value
                                ):  # Прокод по спискам списка
                                    for j, item in enumerate(
                                        sublist
                                    ):  # Проход по значениям из вложенного списка
                                        target_cell_col = start_cell.column + j
                                        target_cell_row = start_cell.row + i
                                        col_letter = get_column_letter(target_cell_col)
                                        self.workbook[location_sheet][
                                            f"{col_letter}{target_cell_row}"
                                        ].value = item
                            else:
                                start_cell.value = value

    def save_to_bytes(self) -> BytesIO:
        if not self.workbook:
            raise ValueError("Excel файл не загружен.")
        output = BytesIO()
        self.workbook.save(output)
        output.seek(0)
        return output

    def save_to_file(self, file_path: str) -> None:
        if self.workbook:
            self.workbook.save(file_path)
        else:
            raise ValueError("Excel файл не загружен.")
